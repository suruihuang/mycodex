"""Reproduce quote-only concise categories; no external paper/structure validation."""
from pathlib import Path
from collections import Counter
from copy import copy
from datetime import datetime, timezone
import csv, hashlib, json, re, time
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent / 'first50_quote_only_mutations.xlsx'
OUTPUT = HERE / 'first50_quote_only_mutations_mechanisms.xlsx'
# Explicit human-readable review decisions keyed by ORIGINAL evidence key.
# category, separate uncertainty/additional interactions, evidence level, rationale
M = {}
def put(keys, category, other, level, reason):
    for key in keys.split():
        M[key] = [category, other, level, reason]
put('E001 E005', 'Not stated', 'Electrostatic is a possible interpretation of the listed charged residues, but is not established by this quote.', 'Unspecified contact', 'Only "interacts" is stated; no salt bridge, hydrogen bond, or mutation effect is identified.')
put('E002 E003 E010 E011 E012 E013 E014 E015', 'electrostatic', 'Electrostatic guidance is an author proposal, not a uniquely demonstrated mechanism; no alternative is specified.', 'Author-proposed mutation mechanism', 'The quote explicitly suggests electrostatic guidance from grouped Arg/Ala and Glu/Ala on-rate changes.')
put('E004', 'packing', 'Possible electrostatic interaction with Asp164/Glu44 is an inference; its physical type is not explicit.', 'Reported contact only', 'Aliphatic Arg64 packing against Trp169 is explicit; no substitution or causal mutation mechanism is reported.')
put('E006 E022', 'electrostatic', 'Not stated', 'Reported contact only', 'A salt bridge is explicitly reported; the category describes the contact, not a demonstrated mutation mechanism.')
put('E007', 'hydrophobic', 'Not stated', 'Reported residue contribution', 'A hydrophobic patch and grouped apparent contributions are explicit, but substitutions and mutation-specific mechanisms are absent.')
put('E008 E016 E018 E020 E038 E045', 'hydrogen bonding', 'Not stated', 'Reported contact only', 'Hydrogen bonds are explicit; no causal mutation mechanism is supplied.')
put('E009', 'Not stated', 'Intermolecular hydrogen bonding is explicitly absent at this site; other physical interaction types are not established.', 'Unspecified contact', 'Hydrophilic clustering alone does not identify the interaction mechanism.')
put('E017 E019 E021', 'packing', 'Electrostatic / hydrogen bonding: polar contacts are removed by the tested replacements but interpreted as relatively unimportant for affinity; no quantitative mechanism shares are stated.', 'Author-proposed mutation mechanism', 'The authors infer dominant intramolecular alkyl packing from R43L/R43M penalties and the contextual alanine comparison; preserved packing is a design rationale, not a verified mutant structure.')
put('E023', 'Not stated', 'No alternative mechanism stated; unchanged affinity is an observed effect, not a molecular mechanism.', 'Effect only', 'S219A affinity is essentially unchanged; a physical explanation is absent.')
put('E024', 'hydrogen bonding + van der Waals', 'Both contact types are stated; their relative contributions are not quantified.', 'Reported contact only', 'Water-mediated hydrogen bonds and van der Waals contacts are explicit; neither is a demonstrated mutation mechanism.')
put('E025 E037 E041 E047', 'hydrogen bonding', 'Not stated', 'Reported contact only', 'Water-mediated hydrogen bonds are explicit; direct hydrogen bonds are also included where stated in this quote.')
put('E026 E032', 'Not stated', 'Physical contact type is not stated; no supported alternative can be selected.', 'Unspecified contact', 'Residue contacts are listed without their physical interaction type or mutation-specific mechanism.')
put('E027', 'Not stated', 'Authors propose weak additional contacts and report partial Glu67 disorder; these do not identify a physical mechanism class.', 'Author interpretation; physical type unspecified', 'No adverse E108A effect and similar His114/Tyr437 effects do not establish hydrogen bonding, electrostatics, or another contact type.')
put('E028 E029 E036', 'Not stated', 'No molecular alternative mechanism is stated.', 'Effect only', 'Binding effects are reported without a physical mechanism or complete replacement identities.')
put('E030 E046', 'van der Waals', 'Hydrophobic burial: small non-polar ASA changes are a co-reported correlate, not a quantified alternative mechanism.', 'Author-linked contacts and mutation effects', 'The quote relates weak affinity changes to van der Waals contacts and small non-polar ASA changes; replacements are unspecified.')
put('E031', 'hydrogen bonding', 'Authors propose loss of the Ser460–His8 interaction as the larger contribution than Ser460–Gln12; mechanism shares are not measured.', 'Author-proposed mutation mechanism', 'This quote itself identifies hydrogen-bonding partners and proposes interaction loss from comparative Ki changes; fold changes are not mechanism probabilities.')
put('E033', 'Not stated', 'Unverified due to interleaved OCR. The Asp41–Lys320 hydrogen bond is a separate contact, not evidence for the R5A mechanism.', 'Corrupted / unresolved evidence', 'The R5A effect text is incomplete; assigning the neighboring hydrogen-bond fragment to R5A would be unsupported.')
put('E034 E039 E040 E042', 'hydrogen bonding + van der Waals', 'Both contact types are stated; their relative contributions are not quantified.', 'Reported contact only', 'The quote explicitly reports both hydrogen bonds and van der Waals contacts, without a tested mutation mechanism.')
put('E035', 'electrostatic', 'Hydrogen bonding is also discussed for the same Lys40–Asp435 contact; these descriptions overlap and do not imply separate additive contributions.', 'Structure and mutagenesis support', 'Two salt bridges are reported, with coupled K40Q/D435A mutation effects supporting the contact assignment; no energetic percentage partition is given.')
put('E043', 'hydrogen bonding', 'The change and substitution are unidentified in this quote; causal attribution remains incomplete.', 'Author-proposed explanation; mutation unresolved', 'Authors explain an unspecified change by loss of interactions and report Arg5–Trp438 hydrogen bonds; do not assign R5A from another row.')
put('E044', 'Not stated', 'No molecular mechanism stated; hydrogen bonding must not be imported from another quote.', 'Effect only', 'The 50-fold dissociation-constant increase establishes weaker binding, not its physical cause.')
put('E048', 'shape complementarity', 'Not stated', 'Reported contact only', 'Interface shape complementarity is explicit; no substitution or energetic mechanism is established.')
put('E049', 'hydrophobic', 'Hydrophobic ASA burial is associated with grouped mutation effects; causal exclusivity and individual contributions are not established.', 'Author-linked contacts and mutation effects', 'The quote associates hydrophobic surface burial with modest affinity effects, without specifying replacements or per-mutant values.')
put('E050', 'van der Waals', 'Packing: the same reported contacts surround Pro88; this is complementary geometry, not an independent percentage contribution.', 'Reported contact only', 'Fifteen van der Waals contacts and enveloping geometry are stated; no mutation mechanism is reported.')

def now(): return datetime.now(timezone.utc).isoformat()
script_start = time.perf_counter()
start = now()
original_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
assert original_hash == 'efc9c19fcf1ececc18ec8ff923ae7b14d8b6a686c1ed7327fe3c4aa31b5da12e'
w = load_workbook(SOURCE)
original = {s.title: list(s.values) for s in w}
rows = original['Row_summary'][1:]
assert len(M) == len(rows) == 50
quote_map = {}
result_map = {}
for row in rows:
    key = row[8]
    decision = M[key]
    if row[3] in quote_map: assert quote_map[row[3]] == decision
    quote_map[row[3]] = decision
    bounds = [int(n) for n in re.findall(r'\d+', str(row[7]))]
    for r in range(bounds[0], bounds[-1] + 1):
        assert r not in result_map
        assert w['Results'].cell(r, 5).value == row[3]
        result_map[r] = decision
assert set(result_map) == set(range(2,122)) and len(quote_map) == 36
new_headers = ['mechanism', 'Uncertain / alternative mechanisms', 'Mechanism evidence level', 'Mechanism classification rationale']
insertions = {'Results': 4, 'Row_summary': 8, 'Evidence_log': 7}
for name, col in insertions.items():
    s = w[name]
    dims = {i: copy(s.column_dimensions[get_column_letter(i)]) for i in range(1, s.max_column+1)}
    s.insert_cols(col, 4)
    for i,d in dims.items():
        target = i+4 if i>=col else i
        d.index = get_column_letter(target)
        s.column_dimensions[d.index] = d
    for i,h in enumerate(new_headers,col):
        c=s.cell(1,i,h);c._style=copy(s.cell(1,col-1)._style)
        s.column_dimensions[get_column_letter(i)].width = [23,65,38,75][i-col]
    for r in range(2,s.max_row+1):
        d = result_map[r] if name == 'Results' else M[rows[r-2][8]]
        for j,v in enumerate(d,col):
            c=s.cell(r,j,v);c._style=copy(s.cell(r,col-1)._style)
            c.alignment=Alignment(wrap_text=True,vertical='top')
            c.fill=PatternFill('solid',fgColor='FFF2CC' if d[0]=='Not stated' else 'E2F0D9')
    s.auto_filter.ref = f'A1:{get_column_letter(s.max_column)}{s.max_row}'
    for row in s:
        for c in row:
            if c.hyperlink: c.hyperlink.ref=c.coordinate

summary = w.create_sheet('Mechanism_summary')
summary.append(['Population','Metric','Category / definition','Count','Denominator','Percentage','Interpretation'])
populations = [('Result entries',list(result_map.values())),('Source rows',[M[r[8]] for r in rows]),('Exact distinct quotes',list(quote_map.values()))]
for pop, values in populations:
    counts=Counter(d[0] for d in values)
    for category,count in sorted(counts.items()):
        summary.append([pop,'Category frequency',category,count,len(values),count/len(values),'One primary label per entry; combined label retained as one category.'])
    for level,count in sorted(Counter(d[2] for d in values).items()):
        summary.append([pop,'Evidence level frequency',level,count,len(values),count/len(values),'Evidence labels distinguish observed contacts from causal mutation explanations.'])
    for label,predicate,meaning in [
        ('Unassigned physical category',lambda d:d[0]=='Not stated','Physical category is unavailable; does not mean there is no mechanism.'),
        ('Uncertainty / additional interaction note',lambda d:d[1]!='Not stated','Any nonempty separate note, including co-reported contacts; not a probability of error.'),
        ('Proposed or unresolved explanation',lambda d:d[2].startswith('Author-proposed') or d[2] in ['Corrupted / unresolved evidence','Author interpretation; physical type unspecified'],'Author proposals and incomplete assignments; overlaps other summary metrics.')]:
        count=sum(predicate(d) for d in values)
        summary.append([pop,'Uncertainty ratio',label,count,len(values),count/len(values),meaning])
for r in range(2,summary.max_row+1):summary.cell(r,6).number_format='0.00%'
notes = [
 ('Dated mechanism copy','2026-09-30 folder follows the user environment date. Recorded activity UTC timestamps use the actual system clock and may have a different date.'),
 ('Original preserved','All original sheet cell values are preserved. Added lowercase mechanism is a concise category alongside the original detailed mechanism. Original workbook is unchanged.'),
 ('Evidence scope','Only the existing quotes and detailed evidence were reviewed. No paper, PDF, PDB record, or external source was newly consulted or validated.'),
 ('Category interpretation','Contact-only rows label the reported interaction, not a demonstrated mutation mechanism. Consult Mechanism evidence level. Mixed hydrogen bonding + van der Waals is a combined category.'),
 ('Percentage definition','Percentages are observed category/evidence/uncertainty frequencies, separately divided by 120 result entries, 50 source rows, and 36 exact distinct quote strings. Duplicates are retained except in exact-distinct-quote summaries.'),
 ('Alternative percentages','Per-row probabilities or energetic percentage shares are Not measured / not estimable from these quotes. No arbitrary confidence percentages were assigned. Frequency ratios are not probabilities that a mechanism is correct.'),
 ('Distinct-quote definition','Exact string equality only; punctuation and small wording variants remain distinct. This is not a count of independent experiments.'),
 ('Activity history','Activity_log A001–A064 is inherited unchanged. Mechanism_activity and activity_log.csv record this new subagent run; durations apply only to explicitly instrumented intervals; exact tokens are Not measured.'),
 ('Source SHA256',original_hash)]
for pair in notes:w['Read_me'].append(pair)
review_end=now()
started=json.loads((HERE/'analysis_started.json').read_text())
activity = [
 ['M001','Not recorded','Not recorded',None,'functions.exec / exec_command / Python openpyxl','Inspect source workbook and inherited logs','120 result entries, 50 source rows and 36 exact quotes; initial inspection preceded instrumentation','Not measured'],
 ['M002',started['timestamp'],review_end,(datetime.fromisoformat(review_end)-datetime.fromisoformat(started['timestamp'])).total_seconds(),'Subagent quote review + exec_command / Python','Review quotes, author decisions, and generate categorized workbook','Instrumented elapsed interval includes review, mapping/script preparation, communication and generation; not pure analysis time','Not measured'],
]
log=w.create_sheet('Mechanism_activity')
log.append(['Event','Started UTC','Finished UTC','Measured elapsed seconds','Tool / process','Action','Outcome / timing limitation','Exact tokens'])
for a in activity:log.append(a)
for s in [summary,log]:
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
    for row in s:
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    for col in range(1,s.max_column+1):s.column_dimensions[get_column_letter(col)].width=35 if col<7 else 90
w.save(OUTPUT)
# Reload to check actual serialized values rather than only the in-memory model.
v=load_workbook(OUTPUT)
for name, oldrows in original.items():
    col=insertions.get(name)
    for ri,row in enumerate(oldrows,1):
        for ci,value in enumerate(row,1):
            assert v[name].cell(ri,ci+4 if col and ci>=col else ci).value == value,(name,ri,ci)
for s in v:
    for row in s:
        for c in row:
            assert c.data_type != 'f'
            if c.hyperlink: assert c.hyperlink.ref==c.coordinate
for pop,_ in populations:
    rr=[r for r in list(v['Mechanism_summary'].values)[1:] if r[0]==pop and r[1]=='Category frequency']
    assert sum(r[3] for r in rr)==rr[0][4] and abs(sum(r[5] for r in rr)-1)<1e-10
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==original_hash
end=now()
validation=['M003',review_end,end,(datetime.fromisoformat(end)-datetime.fromisoformat(review_end)).total_seconds(),'exec_command / Python openpyxl','Save, reload, and validate','All original cell values, source hash, quote assignments, 120-row coverage, no formulas, hyperlink refs and frequency totals passed','Not measured']
activity.append(validation)
log.append(validation)
log.auto_filter.ref=log.dimensions
w.save(OUTPUT)
with (HERE/'activity_log.csv').open('w',newline='',encoding='utf-8-sig') as f:
    writer=csv.writer(f);writer.writerow([c.value for c in log[1]]);writer.writerows(activity)
with (HERE/'mechanism_mapping.json').open('w') as f:json.dump(M,f,indent=2,ensure_ascii=False)
report={'source_sha256':original_hash,'source_rows':50,'result_entries':120,'exact_distinct_quotes':36,'original_cell_values_preserved':True,'source_file_unchanged':True,'formulas':0,'hyperlink_refs_valid':True,'category_percentages_sum_to_100':True,'script_elapsed_seconds':time.perf_counter()-script_start,'script_timing_scope':'Generation and validation only; does not represent total analysis time.','tokens':'Not measured','frequency_counts':{p:dict(Counter(d[0] for d in ds)) for p,ds in populations}}
(HERE/'validation_report.json').write_text(json.dumps(report,indent=2))
(HERE/'README.txt').write_text('\n\n'.join(k+': '+v for k,v in notes)+'\n')
print(json.dumps(report,indent=2))
print(OUTPUT)
