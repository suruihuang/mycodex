from pathlib import Path
from datetime import datetime, timezone
import csv, hashlib, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
BASE=Path('/home/surui-wsl/code/codex/codex-rs/agent_trial_1')
SOURCE=Path('/home/surui-wsl/code/codex/parsed data/quotes_pdf_ids.xlsx')
start=datetime.now(timezone.utc).isoformat()
before=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
src=openpyxl.load_workbook(SOURCE, data_only=False).active
rows=list(src.iter_rows(min_row=2,max_row=51,values_only=True))
D={}
def put(ids, entries):
    for i in ids: D[i]=entries
none='None specified — residue interaction only'
def contact(ids,mechanism,reason): put(ids,[(none,mechanism,reason)])
contact([6,13],'Unspecified','hGH Arg64 contacts receptor Asp164 and Glu44; interaction type is not stated. No substitution is specified.')
put([7,10,42,44,45,47,48,49],[(f'Hormone (identity not stated): {a}{n} → Ala ({c}{n}A)','Electrostatic',f'The quote explicitly names this conversion. Arg-to-Ala conversions have grouped on-rate reductions of about twofold; Glu-to-Ala conversions have grouped on-rate increases of nearly twofold. No individual value or explicit single versus combined construct design is provided. Electrostatic guidance is an author proposal ("suggests"), not a demonstrated causal assignment. Protein identity beyond hormone is not stated.') for a,c,n in [('Arg','R',64),('Arg','R',167),('Arg','R',178),('Glu','E',65),('Glu','E',174)]])
contact([12],'Packing','hGH Arg64 aliphatic side chain packs against Trp169; its guanido group interacts with hGHbp Asp164 and Glu44 without an explicit interaction type. Functional importance is stated; no substitution is given.')
contact([16],'Saltbridge','Receptor linker Glu127 forms an explicitly stated salt bridge to hGH Lys41; no substitution is given.')
contact([25],'Hydrophobic','Ile103, Ile105, Pro106 and Ile165 form a hydrophobic patch with tryptophans; apparent contributions of 1.5–3.5 kcal/mol are grouped. Protein and substitutions are not stated.')
contact([29],'Hydrogenbonding','An intermolecular hydrogen bond connects Lys167 carbonyl oxygen to hGH Asn68 Ne2. The Lys167-bearing protein is not explicitly named; no substitution is given.')
contact([33],'Unspecified','Receptor Arg217/Asn218 cluster with hGH His18/His21/Glu174; intermolecular hydrogen bonds are explicitly absent. Contact or hydrophilicity does not establish a different interaction mechanism.')
contact([52,54,56],'Hydrogenbonding','Arg43 makes stated hydrogen bonds with hGH Asp171 and Thr175. Arg43-bearing protein identity and a substitution are not stated.')
put([53,55,57],[(f'Protein not stated (hGH-binding partner): Arg43 → {target} (R43{code})','Packing',reason+' Author interpretation: intramolecular alkyl packing accounts for most of the energetic contribution; electrostatic interactions with hGH are relatively unimportant for affinity. Quote also explicitly describes a salt bridge, hydrogen bond and van der Waals contacts. Figure 3, Figure 6, Table 2 are cited within the quote; no outside source was inspected.') for target,code,reason in [('Leu','L','R43L is explicitly engineered and analyzed; complex destabilization is 0.5 kcal/mol.'),('Met','M','R43M is explicitly engineered and analyzed; complex destabilization is 1 kcal/mol.'),('Ala','A','The retrospective "alanine substitution" refers to Arg43 in this paragraph (contextual assignment). The quoted contribution exceeds 2.1 kcal/mol, but the quote does not expressly give an exact R43A measurement.')]])
contact([58],'Saltbridge','Arg71 forms a stated salt bridge to hGH Glu56 and buries 80 AÊ 2 upon binding. Protein bearing Arg71 is not stated. OCR unit retained without silently correcting the quote.')
put([65],[('Protein not stated: Ser219 → Ala (S219A)','Unspecified','Affinity is essentially unchanged relative to wild type. No interaction type or causal mechanism is supplied; OCR af®nity/thewild-type retained.')])
contact([75],'Hydrogenbonding','Ang Arg5 and Gln12 form water-mediated hydrogen bonds to hRI Ile436 and Asp435 respectively (Table V). A separate Ang–Asp435 van der Waals contact has an unspecified Ang residue.')
contact([76],'Hydrogenbonding','Direct Trp89–Glu344 hydrogen bond and water-mediated Ang Gly86–RI Glu206/Trp263 and Ang Trp89–RI Glu344 hydrogen bonds are stated (Table V).')
contact([77],'Unspecified','RI Glu344 and Glu401 interact with Ang Trp89, without an interaction type; "both of these RI residues" is unresolved within this quote.')
put([78],[('Ang: Glu108 → Ala (E108A)','Unspecified','No adverse effect is reported for Glu108-to-Ala. Authors suggest additional Tyr437 contacts are weak; partial Glu67 disorder is also stated, but no interaction type or causal mechanism for E108A is established. Table III and Chen and Shapiro, 1997 are quoted references only.'),('Protein not stated: His114 → Not stated','Unspecified','His114 mutation is mentioned, but replacement and its specific consequence are not supplied.'),('Protein not stated: Tyr437 → Not stated','Unspecified','Tyr437 mutation is mentioned, but replacement is not supplied. Similar consequences to His114 mutations motivate an author suggestion that additional contacts to Ang Glu67/Ala106/Glu108 are weak.')])
put([79],[(f'Ang: {r} → Not stated','Unspecified','The replacement is omitted. The three named substitutions are described as producing G values greater than 0.5 kcal/mol, within a group with no binding-energy decrease exceeding 1 kcal/mol. Do not silently convert the printed G to delta-delta-G or assign an individual value; mechanism is absent.') for r in ['His8','Arg32','His114']])
put([80,87],[(f'Ang: {r} → Not stated','Unspecified','Residue modification is stated, but replacement and modification type are not specified. The nine-residue group has no significant effect on RI binding affinity. No interaction type or causal mechanism is supplied.') for r in ['His13','Arg31','Arg33','Lys50','Lys60','Arg66','Asn68','Arg70','Glu108']])
put([81,97],[(f'{p}: {r} → Not stated','Dispersion','Substitutions are stated, but replacements are absent. The group barely alters affinity and forms only van der Waals contacts, with small changes in nonpolar ASA. Dispersion labels the explicitly stated van der Waals interactions; no individual affinity value is assigned.') for p,r in [('hRI','Trp261'),('hRI','Lys320'),('Ang','His84'),('Ang','Glu108')]])
put([82],[('RI: Ser460 deletion (des(460))','Hydrogenbonding','The des(460) derivative has an 8.5-fold increase in Ki. Interpreting des(460) as deletion of the named Ser460 is a notation-based inference. Authors attribute the effect largely to loss of its His8 interaction rather than Gln12.'),('Ang: His8 → Ala (H8A)','Hydrogenbonding','Individually substituting His8 by Ala yields a 4.6-fold change in Ki (metric inferred from the immediately preceding Ki sentence). His8 is a stated hydrogen-bonding partner of RI Ser460.'),('Ang: Gln12 → Ala (Q12A)','Hydrogenbonding','Individually substituting Gln12 by Ala yields a 1.6-fold change in Ki (metric inferred from the immediately preceding Ki sentence). Gln12 is a stated hydrogen-bonding partner of RI Ser460.')])
contact([83],'Unspecified','Ang–RI contact pairs are Arg24–Asn93, Gln93–Tyr150, Asp41–Lys320, Arg121–Asp378, Ser37–Ile459 and Pro38–Ile459. Interaction types and mutations are absent; apparent OCR word "rate" retained.')
put([84],[('Ang: Arg5 → Ala (R5A)','Unspecified','The Arg5-to-Ala mutation is explicit, but its effect is truncated at "also increased". The interleaved Asp41–Lys320 hydrogen-bond description cannot establish the mechanism of R5A. Severe apparent text/OCR interleaving; quote preserved without repair.')])
contact([85],'Hydrogenbonding','Lys40 in the hRI–Ang complex makes two hydrogen bonds to RI Asp435 and van der Waals contacts to Tyr434. Assigning Lys40 to Ang is contextual within this quote. OCR ?-amino retained.')
put([86],[('Ang: Lys40 → Gln (K40Q)','Saltbridge','Lys40-to-Gln weakens binding 1300-fold. The quote reports two structural salt bridges from Lys40 to RI Asp435, following earlier hydrogen-bond hypotheses. No mutant structure is identified.'),('RI: Asp435 → Ala (D435A)','Saltbridge','Asp435-to-Ala decreases affinity almost as much as K40Q; no exact individual fold value is given. The quote describes two Lys40–Asp435 salt bridges.'),('Joint complex: Ang Lys40 → Gln (K40Q) + RI Asp435 → Ala (D435A)','Saltbridge','The joint mutant complex dissociates at the same rate as K40Q with native RI. Preserve this combined experiment without distributing its dissociation result to the component mutations. Two wild-type Lys40–Asp435 salt bridges are stated.')])
contact([88,92],'Hydrogenbonding','Ang Arg5 and Gln12 form water-mediated hydrogen bonds to hRI Ile436 and Asp435 respectively (Table V).')
contact([89],'Hydrogenbonding','Ang identity is not stated for Gln12/His8 in this quote; both hydrogen-bond to hRI Ser460. Gln12 participation in P1 is an author expectation, not a mutation.')
contact([90,91],'Hydrogenbonding','Ang Arg31 forms two hydrogen bonds and van der Waals contacts to RI Asp35; Ang Arg32 forms one hydrogen bond and multiple van der Waals contacts to RI Asp36 (Figure 5E). No substitutions are supplied; OCR ?-helix retained.')
contact([93],'Hydrogenbonding','Arg5 forms two hydrogen bonds to Trp438 carbonyl and van der Waals contacts to Ser439/Glu440 (Figure 5B). Protein identities are not explicitly stated within this quote.')
contact([94],'Hydrogenbonding','Authors explain an unspecified change by loss of interactions with RI; two Arg5–RI Trp438 hydrogen bonds are given. Mutation and the antecedent of "This change" are absent; no mutation is imported from other rows.')
put([95],[('Ang: Arg5 → Ala (R5A)','Unspecified','Dissociation constant increases 50-fold. No interaction type or causal mechanism is stated; dissociation constant is not dissociation rate. Shapiro and Vallee, 1992 is a reference within the quote only.')])
contact([96],'Hydrogenbonding','Two direct hydrogen bonds link hRI Ser289–Ang Gly86 and hRI Glu344–Ang Trp89 respectively.')
contact([98],'Hydrogenbonding','Direct hydrogen bonds connect hRI Ser289–Ang Gly86 and hRI Glu344–Ang Trp89; water-mediated hydrogen bonds connect hRI Trp263–Ang Gly86 and hRI Glu344–Ang Trp89 respectively.')
contact([99],'Complementarity','High shape complementarity is explicitly stated for hRI–Ang surfaces, especially hRI 434–438/460 with the Ang active site and hRI Trp261/263/318 with the Ang Trp89 loop. No mutation or mutation-induced effect is stated.')
put([100],[(f'hRI: Trp{n} → Not stated','Hydrophobic','Replacements are mentioned but their identities are not stated. The group has printed G values of 1.0–1.5 kcal/mol and reductions of 34–60 AÊ 2 in hydrophobic ASA on complex formation. This supports hydrophobic burial as an author-associated interaction; no individual values or corrected OCR metric are assigned.') for n in [263,318,375]])
contact([101],'Dispersion','Trp318 and Trp375 envelop Ang Pro88 with 15 total van der Waals contacts. Protein bearing the tryptophans is not explicitly stated; Dispersion labels van der Waals interactions. No substitution is given.')
assert set(D)=={r[0] for r in rows}
w=openpyxl.Workbook(); result=w.active; result.title='Results'
result.append(['PDB','Mutation DECODED','Mechanism','Citation','Quote','Reason'])
summary=w.create_sheet('Source summary'); summary.append(['Data row','Source Excel row','Source ID','Citation label','Quote','Result Excel rows','Entry count','Limits'])
ev=w.create_sheet('Evidence'); ev.append(['Source ID','Source Excel row','Result Excel row','Evidence source','Quote','Assignment','Mechanism','Evidence and caveats'])
for ix,(sid,cite,quote) in enumerate(rows,1):
    rr=[]
    for mutation,mechanism,reason in D[sid]:
        ri=result.max_row+1; rr.append(str(ri))
        reason+=' PDB not stated in the quote; citation filename is not structure evidence. Numbering is quote numbering; chains and deposited-mutant status are not verified.'
        result.append(['Not stated',mutation,mechanism,f'{cite} (source label only; supplied quote, Excel row {ix+1}, ID {sid})',quote,reason])
        ev.append([str(sid),str(ix+1),str(ri),'This row quote only',quote,mutation,mechanism,reason])
    summary.append([str(ix),str(ix+1),str(sid),cite,quote,', '.join(rr),str(len(rr)),'No external verification; no cross-row enrichment; protein abbreviations retained as quoted.'])
activity=[['UTC','Action','Details','Exact tokens'],[start,'Setup','Read applicable skill; user restrictions override requests for external verification. Read only header and first 50 data rows from original workbook. Independently authored interpretations and this script.','Not measured'],[datetime.now(timezone.utc).isoformat(),'Extraction','Quotes are sole scientific evidence. Citation labels preserved only. No web, papers, previous results/scripts, other-agent outputs, or cross-row scientific enrichment used. Identical quotations use identical interpretations. No package installation or background scientific work.','Not measured'],[datetime.now(timezone.utc).isoformat(),'Caveats','One-word mechanisms summarize explicit interaction types; effects alone use Unspecified. Dispersion denotes quoted van der Waals contacts. Ambiguous construct grouping, protein identity, OCR, deletion notation and contextual assignments are flagged. No PDB identifier appears in supplied quotes.','Not measured']]
act=w.create_sheet('Activity')
for r in activity:act.append(r)
for ws in w:
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476B')
    for row in ws:
        for cell in row:
            cell.value=str(cell.value) if cell.value is not None else ''
            cell.data_type='s';cell.alignment=Alignment(vertical='top',wrap_text=True)
    for col in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width=24
    for col in (['E','F'] if ws.title=='Results' else ['E','H'] if ws.title=='Evidence' else ['E','H'] if ws.title=='Source summary' else ['C']):ws.column_dimensions[col].width=85
w.save(BASE/'results.xlsx')
check=openpyxl.load_workbook(BASE/'results.xlsx',data_only=False)
assert check['Source summary'].max_row==51
for row in check['Source summary'].iter_rows(min_row=2,values_only=True):
    quote=rows[int(row[0])-1][2];assert row[4]==quote
    for ri in row[5].split(', '):assert check['Results'].cell(int(ri),5).value==quote
assert all(len(r[2].split())==1 for r in check['Results'].iter_rows(min_row=2,values_only=True))
assert not any(c.data_type=='f' for ws in check for row in ws for c in row)
assert before==hashlib.sha256(SOURCE.read_bytes()).hexdigest()
validation=['50 source rows covered','All expanded quotes equal their source quote','All mechanisms one word','No spreadsheet formulas','Original SHA256 unchanged']
activity.append([datetime.now(timezone.utc).isoformat(),'Validation','; '.join(validation)+f'; {result.max_row-1} result entries.','Not measured'])
act.append(activity[-1])
for c in act[act.max_row]:c.data_type='s';c.alignment=Alignment(wrap_text=True,vertical='top')
act.auto_filter.ref=act.dimensions
w.save(BASE/'results.xlsx')
for name, sheet in [('evidence_log.csv',ev),('activity_log.csv',act)]:
    with (BASE/name).open('w',newline='',encoding='utf-8-sig') as f:csv.writer(f).writerows(sheet.values)
report={'source_rows':50,'result_entries':result.max_row-1,'evidence_entries':ev.max_row-1,'source_sha256':before,'checks':validation,'exact_tokens':'Not measured'}
(BASE/'validation.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report,indent=2))
